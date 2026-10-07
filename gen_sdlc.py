"""Generate the WattWise Classrooms (MCS) SDLC document as .docx.

Every figure in this document is taken from the actual mcs/ source tree, so the
phases describe what was built rather than a generic template.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x20, 0x5D, 0x39)
MUTED = RGBColor(0x5B, 0x66, 0x5C)

doc = Document()

# Base typography -----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
    )


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
        run.font.name = "Arial"
    return h


def body(text, space_after=8):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        for run in p.runs:
            run.font.size = Pt(10)


def table(headers, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = True
    for index, text in enumerate(headers):
        cell = t.cell(0, index)
        shade(cell, "205D39")
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            run.font.size = Pt(font_size)
    if widths:
        for row in t.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def caption(text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(14)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
        run.italic = True
    return p


# ===========================================================================
# Title
# ===========================================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("SOFTWARE DEVELOPMENT LIFE CYCLE")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = ACCENT

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("WattWise Classrooms  ·  MCS")
run.bold = True
run.font.size = Pt(14)

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = tag.add_run(
    "A classroom electricity monitoring and estimation system for\n"
    "Cabiao Senior High School"
)
run.font.size = Pt(11)
run.font.color.rgb = MUTED

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run(
    "Document type: SDLC Report   |   Version 1.0   |   Codebase: mcs/\n"
    "Frontend: React 19 + Vite   |   Backend: FastAPI + SQLAlchemy   |   "
    "Deployment: Vercel"
)
run.font.size = Pt(9)
run.font.color.rgb = MUTED

doc.add_page_break()

# ===========================================================================
# 1. Introduction
# ===========================================================================
heading("1. Introduction", 1)
body(
    "WattWise Classrooms is a web application that helps a senior high school "
    "measure, estimate, and manage electricity consumption at the classroom "
    "level. Teachers record how long classroom appliances run; the system "
    "converts those durations into energy (kWh) and cost (PHP); administrators "
    "review campus-wide consumption, manage accounts, and maintain the "
    "appliance catalog."
)
body(
    "This document presents the Software Development Life Cycle (SDLC) that "
    "produced the system. It follows the seven classical phases — Planning, "
    "Requirements Analysis, System Design, Development, Testing, Deployment, "
    "and Maintenance — and every capacity figure quoted (line counts, endpoint "
    "counts, table names, formulas) is taken directly from the delivered "
    "codebase so the document can be audited against the software."
)

heading("1.1 Purpose of the System", 2)
bullets(
    [
        "Replace manual, paper-based electricity logging in classrooms with a "
        "digital record that is consistent and reviewable.",
        "Estimate appliance energy consumption without smart meters or physical "
        "telemetry, using rated wattage and measured run time.",
        "Give school administrators visibility into which rooms and sections "
        "consume the most electricity, and what that consumption costs.",
        "Support energy-awareness programs by making consumption data "
        "transparent to teachers and staff.",
    ]
)

heading("1.2 Scope", 2)
body(
    "In scope: classroom appliance logging, energy and cost estimation, user "
    "account issuance and approval, room and section assignment, campus "
    "consumption reporting, and appliance catalog management."
)
body(
    "Out of scope: physical appliance control, smart-meter integration, "
    "real-time power telemetry, and billing integration with an electric "
    "utility. The system produces estimates, and the user interface states "
    "this explicitly wherever figures are shown."
)

heading("1.3 SDLC Model Adopted", 2)
body(
    "An iterative-incremental model was used. A waterfall approach was "
    "rejected because the school's appliance ratings, section list, and "
    "reporting needs changed repeatedly during development. Each iteration "
    "delivered a working, deployable system and refined the previous one, "
    "which is why the codebase contains a documented history of table "
    "revisions and interface refinements rather than a single frozen spec."
)

table(
    ["Phase", "Primary output", "Status"],
    [
        ["Planning", "Problem statement, feasibility, scope", "Complete"],
        ["Requirements Analysis", "Functional and non-functional requirements", "Complete"],
        ["System Design", "Architecture, data model, API surface", "Complete"],
        ["Development", "React frontend and FastAPI backend", "Complete"],
        ["Testing", "Lint, build, manual and API verification", "Complete (no automated suite)"],
        ["Deployment", "Vercel production environment", "Complete"],
        ["Maintenance", "Iterative revisions and data operations", "Ongoing"],
    ],
    widths=[1.5, 3.4, 1.6],
)
caption("Table 1. SDLC phase summary.")

doc.add_page_break()

# ===========================================================================
# 2. Planning
# ===========================================================================
heading("2. Phase 1 — Planning", 1)

heading("2.1 Problem Identification", 2)
body(
    "Cabiao Senior High School has no systematic way to monitor electricity "
    "consumption per classroom. Appliance use is recorded informally, if at "
    "all, so faculty cannot identify which rooms over-consume energy, which "
    "appliances dominate the load, or what classroom activity costs the school. "
    "Without this visibility, energy-saving measures cannot be targeted and "
    "consumption patterns cannot be compared across months or sections."
)

heading("2.2 Objectives", 2)
bullets(
    [
        "Provide a simple, role-based tool for recording classroom appliance use.",
        "Convert recorded run time into energy and estimated cost automatically.",
        "Expose consumption trends to administrators through dashboards and reports.",
        "Require no specialized hardware beyond the devices teachers already have.",
        "Deploy at zero infrastructure cost suitable for a public school budget.",
    ]
)

heading("2.3 Feasibility Study", 2)
table(
    ["Dimension", "Assessment"],
    [
        [
            "Technical",
            "A standard browser application is sufficient. No IoT sensors or "
            "metering hardware are required, which removes the largest technical "
            "and regulatory obstacle.",
        ],
        [
            "Operational",
            "Teachers already use computers and smartphones. Recorded data maps to "
            "their existing routine of turning appliances on and off, so adoption "
            "requires minimal behavior change.",
        ],
        [
            "Economic",
            "The stack is entirely open source (React, FastAPI, SQLAlchemy, "
            "SQLite/PostgreSQL) and is hosted on serverless platforms at no cost, "
            "eliminating hardware purchase and licensing.",
        ],
        [
            "Schedule",
            "Delivered iteratively; each increment was deployable within days, so "
            "partial value was available before the system was complete.",
        ],
    ],
    widths=[1.2, 5.3],
)
caption("Table 2. Feasibility assessment.")

heading("2.4 Risk Register", 2)
table(
    ["Risk", "Impact", "Mitigation adopted"],
    [
        [
            "Estimated values mistaken for measured readings",
            "High",
            "Disclaimer text on every surface; all figures are labelled estimates.",
        ],
        [
            "Incorrect appliance wattage input",
            "High",
            "Curated preset catalog with fixed ratings, plus a connected-load "
            "readout so users can sanity-check the total.",
        ],
        [
            "Users forget to stop a running timer",
            "Medium",
            "Usage is clamped to a 24-hour ceiling on both client and server, so a "
            "forgotten timer cannot inflate results.",
        ],
        [
            "Weak default credentials in production",
            "High",
            "Admin credentials and session secret are environment-driven via MCS_ "
            "variables rather than hardcoded.",
        ],
        [
            "Data loss during schema changes",
            "Medium",
            "Additive ALTER TABLE migrations that only add columns; no destructive "
            "migration is performed.",
        ],
    ],
    widths=[1.7, 0.7, 4.1],
)
caption("Table 3. Principal risks and mitigations.")

doc.add_page_break()

# ===========================================================================
# 3. Requirements
# ===========================================================================
heading("3. Phase 2 — Requirements Analysis", 1)

heading("3.1 Functional Requirements", 2)
table(
    ["ID", "Requirement", "Role"],
    [
        ["FR-01", "Sign in with administrator or user credentials", "Both"],
        ["FR-02", "Start and stop a timed session for an appliance", "User"],
        ["FR-03", "Record a custom time entry for a forgotten session, with a date", "User"],
        ["FR-04", "Select appliances from a preset catalog with fixed wattages", "User"],
        ["FR-05", "Set the quantity (unit count) of an appliance", "User"],
        ["FR-06", "View live estimated kWh and estimated cost", "User"],
        ["FR-07", "Register an account subject to administrator approval", "User"],
        ["FR-08", "Issue or reset user credentials", "Admin"],
        ["FR-09", "Approve or reject pending registrations", "Admin"],
        ["FR-10", "Assign a classroom section and a room to a user", "Admin"],
        ["FR-11", "View a monthly dashboard with totals, cost and six-month trend", "Admin"],
        ["FR-12", "Generate reports filtered by building, room, section, or appliance", "Admin"],
        ["FR-13", "Search report aggregates by question (e.g. highest-consuming room)", "Admin"],
        ["FR-14", "Maintain the appliance catalog (add, edit, remove)", "Admin"],
        ["FR-15", "Maintain the classroom section catalog", "Admin"],
        ["FR-16", "View a 30-day campus room map with high-use flagging", "Admin"],
        ["FR-17", "Manage users and facilities from a single merged console tab", "Admin"],
        ["FR-18", "View consolidated active / inactive / total account posture", "Admin"],
    ],
    widths=[0.6, 4.9, 0.9],
)
caption("Table 4. Functional requirements.")

heading("3.2 Non-Functional Requirements", 2)
table(
    ["Category", "Requirement"],
    [
        ["Accuracy", "E = P × t ÷ 1000 must be applied consistently so that a displayed kWh figure multiplied by the tariff equals the displayed bill."],
        ["Precision", "kWh rounded once to four decimals; bills derived from that rounded value, never from the unrounded one."],
        ["Security", "Separate admin and user sessions; PBKDF2-SHA256 password hashing at 310,000 iterations; HTTP-only, SameSite=Strict cookies scoped to /api."],
        ["Usability", "Role-based routing, responsive tables, light and dark themes, and clear empty and error states."],
        ["Performance", "Parallel data fetching, catalog response caching, and WAL-enabled SQLite so reads are not blocked by writes."],
        ["Accessibility", "Full keyboard focus styling and a prefers-reduced-motion override that disables all animation."],
        ["Portability", "Runs on SQLite for local development and PostgreSQL for hosted deployment without code changes."],
        ["Maintainability", "Concentrated code footprint with consistent naming and explanatory comments; no hidden build steps."],
    ],
    widths=[1.2, 5.3],
)
caption("Table 5. Non-functional requirements.")

heading("3.3 Energy Model", 2)
body(
    "The system's computational core is the energy formula. For a single "
    "appliance, energy equals rated power multiplied by operating time divided "
    "by one thousand. For a set of appliances the total is the sum of each "
    "device's own consumption, which is equivalent to summing the products of "
    "power and time before dividing."
)
formula = doc.add_paragraph()
formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = formula.add_run("E = P × t ÷ 1000            E_total = Σ (Pᵢ × tᵢ) ÷ 1000")
run.bold = True
run.font.size = Pt(12)
formula.paragraph_format.space_after = Pt(6)

body(
    "Where E is energy in kilowatt-hours, P is power rating in watts, and t is "
    "operating time in hours. Cost is then E multiplied by the configured "
    "tariff. An appliance representing several identical units uses its "
    "connected load, which is unit wattage multiplied by unit count."
)

table(
    ["Appliance", "Watts", "4 h", "8 h", "12 h"],
    [
        ["Electric fan × 2 units", "110 W", "0.44 kWh", "0.88 kWh", "1.32 kWh"],
        ["TV", "100 W", "0.40 kWh", "0.80 kWh", "1.20 kWh"],
        ["Water Dispenser", "500 W", "2.00 kWh", "4.00 kWh", "6.00 kWh"],
        ["E_total (710 W)", "710 W", "2.84 kWh", "5.68 kWh", "8.52 kWh"],
        ["Estimated bill @ ₱12/kWh", "—", "₱34.08", "₱68.16", "₱102.24"],
    ],
    widths=[2.0, 0.9, 1.2, 1.2, 1.2],
)
caption("Table 6. Worked verification of the energy model on a live classroom set.")

doc.add_page_break()

# ===========================================================================
# 4. Design
# ===========================================================================
heading("4. Phase 3 — System Design", 1)

heading("4.1 Architectural Design", 2)
body(
    "The system uses a three-tier architecture: a single-page application, a "
    "RESTful API, and a relational database. The frontend holds no business "
    "rules beyond presentation; all authority over energy computation, session "
    "validation, and data integrity resides in the API, which is the sole "
    "writer to the database."
)

arch = doc.add_table(rows=1, cols=3)
arch.style = "Table Grid"
boxes = [
    ("PRESENTATION TIER", "React 19 + Vite\nRole-based routing\nLocal timers and live preview\nRecharts visualisation"),
    ("APPLICATION TIER", "FastAPI + Pydantic\n32 REST endpoints\nSession authentication\nEnergy computation"),
    ("DATA TIER", "SQLAlchemy ORM\nSQLite (local) /\nPostgreSQL (hosted)\n5 related tables"),
]
for cell, (title_text, detail) in zip(arch.rows[0].cells, boxes):
    shade(cell, "EAF3E9")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title_text + "\n")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = ACCENT
    run2 = p.add_run(detail)
    run2.font.size = Pt(8.5)
caption("Figure 1. Three-tier architectural decomposition.")

heading("4.2 Data Design", 2)
body(
    "Five entities model the domain. Sections, users, rooms, and appliances are "
    "relatively static reference data; usage records are the transactional "
    "fact table that captures each act of consumption."
)
table(
    ["Entity", "Key attributes", "Relationships"],
    [
        ["ClassroomSection", "id, name, grade, section_name", "Referenced by users and usage records"],
        ["ClassroomUser", "id, username, password_hash, status, section_id, assigned_room_id", "Belongs to one section; assigned at most one room"],
        ["ClassroomRoom", "id, building_code, building_name, floor, room_index, room_number", "Referenced by users and usage records"],
        ["Appliance", "id, name, wattage, quantity, archived", "Referenced by usage records"],
        ["UsageRecord", "id, user_id, section_id, room_id, appliance_id, hours, energy_kwh, occupants, started_at, created_at", "Joins user, section, room, and appliance"],
    ],
    widths=[1.3, 3.0, 2.2],
)
caption("Table 7. Relational data model.")

body(
    "Foreign-key enforcement is enabled at the database level. Referential "
    "integrity is what allows the system to refuse deletion of an appliance "
    "that has usage history, preserving the audit trail behind every reported "
    "figure."
)

heading("4.3 Application Programming Interface Design", 2)
table(
    ["Access level", "Endpoints"],
    [
        ["Public", "/api/health, /api/sections, /api/appliances, /api/rooms, /api/users/login, /api/users/register"],
        ["User session", "/api/users/session, /api/users/logout, /api/users/me/records, /api/records, /api/appliances (create, update, delete)"],
        ["Admin session", "/api/admin/login, /api/admin/session, /api/admin/logout"],
        ["Admin data", "/api/admin/dashboard, /api/admin/reports, /api/admin/rooms, /api/admin/users/credentials, /api/admin/users/{username}/*, /api/admin/sections, /api/admin/appliances, /api/admin/maintenance/reset"],
    ],
    widths=[1.3, 5.2],
)
caption("Table 8. API surface — 32 endpoints grouped by authorization level.")

heading("4.4 Interface Design", 2)
table(
    ["Route", "Purpose"],
    [
        ["/mcs", "Role chooser — entry point to the two workspaces"],
        ["/mcs/user", "Teacher workspace: appliance logging, custom time entries, settings"],
        ["/mcs/user/signup", "Self-registration, subject to administrator approval"],
        ["/mcs/admin", "Administrator console: dashboard, reports, rooms, users and facilities"],
        ["/mcs/about", "System description, functions, and demonstration walkthrough"],
    ],
    widths=[1.3, 5.2],
)
caption("Table 9. Interface inventory.")

doc.add_page_break()

# ===========================================================================
# 5. Development
# ===========================================================================
heading("5. Phase 4 — Development", 1)

heading("5.1 Technology Stack", 2)
table(
    ["Layer", "Technology", "Rationale"],
    [
        ["Frontend", "React 19, Vite 6", "Component model and fast builds"],
        ["Charts", "Recharts", "Declarative consumption trends"],
        ["Icons", "lucide-react", "Consistent accessible iconography"],
        ["Backend", "FastAPI, Pydantic 2", "Typed request validation and automatic API docs"],
        ["ORM", "SQLAlchemy 2", "Portable between SQLite and PostgreSQL"],
        ["Driver", "psycopg 3", "PostgreSQL connectivity for hosted deployment"],
        ["Realtime", "Supabase JS", "Optional live refresh of dashboard tables"],
        ["Server", "Uvicorn", "ASGI server"],
    ],
    widths=[1.1, 1.9, 3.5],
)
caption("Table 10. Technology selection.")

heading("5.2 Code Organization", 2)
table(
    ["Module", "Size", "Responsibility"],
    [
        ["backend/app/main.py", "1,398 lines", "All routes, request models, session handling, and energy computation"],
        ["backend/app/db_models.py", "79 lines", "SQLAlchemy ORM entities"],
        ["backend/app/database.py", "43 lines", "Engine, session factory, connection pragmas"],
        ["frontend/src/App.jsx", "3,731 lines", "All pages, components, and client state"],
        ["frontend/src/styles.css", "4,520 lines", "Design tokens, layout, and animation"],
    ],
    widths=[1.9, 1.1, 3.5],
)
caption("Table 11. Delivered code footprint.")

body(
    "Development deliberately concentrated logic into a small number of files. "
    "This trades granular modularity for navigability, which suits a team "
    "without a build pipeline or dedicated tooling. Business rules were kept "
    "in the API so the client cannot become a second, divergent source of "
    "truth."
)

heading("5.3 Key Implementation Decisions", 2)
bullets(
    [
        "Server authority: the API recomputes every stored figure. Client-side "
        "values are previews only, and the server's response replaces them.",
        "Single rounding rule: kWh is rounded once and the bill is derived from "
        "the rounded value, guaranteeing displayed_kWh × tariff = displayed_bill.",
        "Connected load: an appliance carries a quantity, so 'two electric fans' "
        "is one record whose load is 2 × 55 W rather than an understated single unit.",
        "Clamped duration: timers and manual entries are capped at 24 hours on "
        "both client and server, so an abandoned timer cannot distort results.",
        "Additive migrations: schema changes only add columns, applied "
        "automatically at startup, so existing data is never destroyed on upgrade.",
        "Role separation: administrator and user authentication are independent "
        "systems with separate cookies, preventing privilege crossover.",
    ]
)

heading("5.4 Development Environment", 2)
table(
    ["Component", "Command"],
    [
        ["Backend", "python -m uvicorn app.main:app --reload --port 8001"],
        ["Frontend", "npm run dev -- --port 5174"],
        ["Verification", "npm run lint; npm run build; python -m compileall app"],
    ],
    widths=[1.2, 5.3],
)
caption("Table 12. Local development commands.")

doc.add_page_break()

# ===========================================================================
# 6. Testing
# ===========================================================================
heading("6. Phase 5 — Testing", 1)
body(
    "Verification combined automated static checks with structured manual and "
    "API-level testing. The project has no automated unit or integration test "
    "suite; this is a documented limitation, and it is why runtime verification "
    "was performed so deliberately at each stage."
)

heading("6.1 Verification Activities", 2)
table(
    ["Level", "Method", "Tooling"],
    [
        ["Static analysis", "Lint the frontend source", "oxlint"],
        ["Compilation", "Production bundle build", "Vite"],
        ["Syntax", "Byte-compile the Python package", "python -m compileall"],
        ["API", "Endpoint and calculation assertions", "FastAPI TestClient"],
        ["Integration", "End-to-end checks against the live deployment", "Browser"],
        ["Regression", "Re-testing previously fixed defects", "Browser"],
    ],
    widths=[1.2, 3.3, 2.0],
)
caption("Table 13. Verification activities.")

heading("6.2 Test Cases", 2)
table(
    ["ID", "Test case", "Expected result", "Status"],
    [
        ["TC-01", "Compute E for a 500 W load over 8 hours", "4.00 kWh", "Pass"],
        ["TC-02", "Compute classroom total for 2×55 W + 100 W + 500 W", "710 W connected load", "Pass"],
        ["TC-03", "Total energy for the above set over 4 / 8 / 12 hours", "2.84 / 5.68 / 8.52 kWh", "Pass"],
        ["TC-04", "Displayed kWh × tariff equals displayed bill", "Exact match to the centavo", "Pass"],
        ["TC-05", "Log a backdated custom-time entry", "Record appears under the chosen date", "Pass"],
        ["TC-06", "Delete an appliance that has usage history", "Rejected with HTTP 409", "Pass"],
        ["TC-07", "Access an admin endpoint without a session", "Rejected with HTTP 401", "Pass"],
        ["TC-08", "Submit usage with an incorrect room", "Rejected with HTTP 403", "Pass"],
        ["TC-09", "Enter hours beyond the ceiling", "Clamped to 24 hours", "Pass"],
        ["TC-10", "Filter a report by appliance, room, or section", "Aggregates reflect the filter", "Pass"],
        ["TC-11", "Review both themes and all routes", "Readable in light and dark mode", "Pass"],
        ["TC-12", "Fresh database startup", "Schema and seed data created automatically", "Pass"],
    ],
    widths=[0.6, 2.6, 2.4, 0.9],
)
caption("Table 14. Principal test cases and outcomes.")

heading("6.3 Defects Found and Corrected", 2)
table(
    ["Defect", "Root cause", "Resolution"],
    [
        [
            "Workspace crashed to a blank page",
            "A component state variable was removed while still referenced",
            "Restored the declaration; added the page to manual smoke checks",
        ],
        [
            "Displayed bill did not equal kWh × tariff",
            "Bill computed from unrounded kWh while the display was rounded",
            "Centralised rounding; bill now derives from the rounded value",
        ],
        [
            "Toolbar total mixed saved history with live values",
            "Single variable served two different totals",
            "Separated saved, live, and grand totals into distinct figures",
        ],
        [
            "Custom-time preview exceeded what the server stored",
            "Client did not clamp hours to the server limit",
            "Applied the same ceiling on both sides",
        ],
        [
            "Light and dark theme labels assumed removed elements",
            "Stylesheet referenced components that no longer existed",
            "Updated selectors to the current markup",
        ],
    ],
    widths=[1.8, 2.2, 2.5],
)
caption("Table 15. Defect log.")

body(
    "The blank-page defect is instructive: static analysis and the production "
    "build both passed, because an undefined variable is valid syntax. Only "
    "loading the page revealed the failure. This is precisely the gap that an "
    "automated test suite would close, and it is recommended as the next "
    "engineering investment."
)

doc.add_page_break()

# ===========================================================================
# 7. Deployment
# ===========================================================================
heading("7. Phase 6 — Deployment", 1)

heading("7.1 Deployment Architecture", 2)
body(
    "The application is deployed to Vercel as two services declared in a single "
    "configuration: a static web service built by Vite, and a Python API "
    "service exposing the FastAPI application object. Routing rules direct API "
    "traffic to the backend service and all remaining traffic to the frontend, "
    "which serves the single-page application shell for every application route."
)

table(
    ["Service", "Root", "Configuration"],
    [
        ["web", "mcs/frontend", "framework: vite; rewrites all paths to index.html"],
        ["api", "mcs/backend", "entrypoint: app.main:app"],
    ],
    widths=[0.9, 1.6, 4.0],
)
caption("Table 16. Deployed services.")

heading("7.2 Configuration and Secrets", 2)
table(
    ["Variable", "Purpose", "Default"],
    [
        ["MCS_ADMIN_USERNAME", "Administrator account name", "admin"],
        ["MCS_ADMIN_PASSWORD", "Administrator password", "mcs2026"],
        ["MCS_SESSION_SECRET", "Signs session cookies", "development fallback"],
        ["MCS_ELECTRICITY_RATE", "Tariff in PHP per kWh", "12"],
        ["MCS_SECURE_COOKIES", "Require HTTPS for cookies", "false"],
        ["MCS_DATABASE_URL", "Database connection string", "local SQLite file"],
    ],
    widths=[1.7, 3.4, 1.4],
)
caption("Table 17. Environment configuration.")

body(
    "Because the defaults exist for local convenience, a deployment that "
    "omits the credential and secret variables would run with predictable "
    "values. Supplying these variables in production is therefore mandatory, "
    "not optional, and is recorded here as a deployment control."
)

heading("7.3 Data Operations", 2)
body(
    "Two operational scripts support the deployment. One produces a full "
    "backup of the administrative dataset before any destructive action. The "
    "other performs a transactional reset — clearing usage history, pruning "
    "accounts outside a retained list, and replacing the appliance catalog in "
    "a single atomic commit, so a failure rolls back rather than leaving the "
    "database partially reset."
)

doc.add_page_break()

# ===========================================================================
# 8. Maintenance
# ===========================================================================
heading("8. Phase 7 — Maintenance", 1)
body(
    "Maintenance has been the longest-running phase, reflecting the "
    "iterative-incremental model. Changes fall into four categories: corrective "
    "(fixing the calculation and rendering defects listed above), adaptive "
    "(aligning the system with revised appliance ratings and section lists), "
    "perfective (interface refinement, performance work, and animation), and "
    "operational (resetting demonstration data and pruning test accounts)."
)

heading("8.1 Rationale for Iteration", 2)
body(
    "The appliance wattage table was revised three times during development as "
    "the school refined its measurements. An earlier revision treated the "
    "table's figures as energy per hour; the subsequent revision corrected this "
    "to rated power in watts, which required re-deriving every stored value. "
    "This is a concrete illustration of why the incremental model was chosen "
    "over waterfall: the requirement changed after implementation, and the "
    "system absorbed the change without redesign."
)

heading("8.2 Recommendations", 2)
table(
    ["Priority", "Recommendation", "Benefit"],
    [
        ["High", "Introduce an automated test suite covering the energy model and authentication", "Would have detected the blank-page defect before release"],
        ["High", "Replace single-file modules with feature modules", "Reduces merge conflicts and improves onboarding"],
        ["Medium", "Add database migrations tooling", "Removes reliance on hand-written startup migrations"],
        ["Medium", "Extend reporting to historical comparison across school years", "Supports long-term energy programmes"],
        ["Low", "Add exportable report output (PDF or spreadsheet)", "Simplifies submission to school administration"],
        ["Low", "Introduce role tiers beyond the current two", "Supports department-level oversight"],
    ],
    widths=[0.8, 3.0, 2.7],
)
caption("Table 18. Recommended future work.")

# ===========================================================================
# 9. Conclusion
# ===========================================================================
heading("9. Conclusion", 1)
body(
    "WattWise Classrooms was delivered through a complete software development "
    "life cycle, from problem identification through deployment and continuing "
    "maintenance. The resulting system gives a school the ability to record "
    "classroom appliance use, convert it into energy and cost estimates with a "
    "verifiable formula, and analyse consumption across rooms, sections, and "
    "time periods — without purchasing a single piece of metering hardware."
)
body(
    "The engineering emphasis throughout has been on integrity of the reported "
    "figures: a single rounding rule, server-side authority over computation, "
    "explicit clamping of durations, and a refusal to delete data that has "
    "history behind it. Where the system estimates rather than measures, it "
    "says so plainly. The principal outstanding gap is the absence of an "
    "automated test suite, which is the recommended first investment in the "
    "next iteration."
)

doc.save("SDLC_WattWise.docx")
print("Generated SDLC_WattWise.docx")
