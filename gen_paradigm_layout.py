import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_box(doc, title, body, fill="4A90E2"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(2.2)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title + "\n")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(255, 255, 255)
    run2 = p.add_run(body)
    run2.font.size = Pt(7)
    run2.font.color.rgb = RGBColor(255, 255, 255)
    return tbl

doc = docx.Document()
doc.add_paragraph('Figure 2: Research Paradigm').bold = True
p = doc.add_paragraph('RESEARCH PARADIGM')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold = True

# Create a 2-column table to mimic the zigzag layout
table = doc.add_table(rows=5, cols=2)
table.autofit = False
for col in table.columns:
    col.width = Inches(3)

# Data for boxes
# (Row, Col, Title, Body)
boxes = [
    (0, 0, "Problem Identification", "Overconsumption of energy in Cabiao Senior High School and lack of managing and monitoring of electricity."),
    (0, 1, "Planning", "Selection of software (Visual Studio Code and OpenCode) and the model's features needed for developing."),
    (1, 0, "Data Harvesting", "Collect relevant classroom electricity data, such as operating hours, appliances used, power consumption, and kWh usage."),
    (1, 1, "Development", "Process and organize the collected data, then use it to train and develop the WattWise."),
    (2, 0, "Testing", "Test the model in predicting classroom electricity consumption, reliability in monitoring kWh usage, and consistency of its predictions."),
    (2, 1, "Data Analysis", "Examining the collected results and compare the predicted electricity consumption with actual consumption."),
    (3, 0, "Results", "Present the model's accuracy and electricity consumption patterns, showing how WattWise can support schools in better energy monitoring and management.")
]

for row, col, title, body in boxes:
    cell = table.cell(row, col)
    set_cell_background(cell, "4A90E2")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title + "\n")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(255, 255, 255)
    run2 = p.add_run(body)
    run2.font.size = Pt(7)
    run2.font.color.rgb = RGBColor(255, 255, 255)

doc.add_paragraph('Research Paradigm').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save('Research_Paradigm_Layout.docx')
