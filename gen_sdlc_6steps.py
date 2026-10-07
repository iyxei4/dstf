"""Generate a simplified 6-step SDLC document for WattWise Classrooms (mcs/)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x20, 0x5D, 0x39)
MUTED = RGBColor(0x5B, 0x66, 0x5C)

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
    )


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
        run.font.name = "Arial"
    return h


def body(text, after=8):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    return p


def caption(text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(14)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED
        run.italic = True


def table(headers, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = t.cell(0, index)
        shade(cell, "205D39")
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = t.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            p = cells[index].paragraphs[0]
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
    if widths:
        for row in t.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ---------------------------------------------------------------- Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("SOFTWARE DEVELOPMENT LIFE CYCLE")
run.bold = True
run.font.size = Pt(19)
run.font.color.rgb = ACCENT

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("WattWise Classrooms  ·  6-Step SDLC")
run.bold = True
run.font.size = Pt(13)

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = tag.add_run("Classroom electricity monitoring and estimation system")
run.font.size = Pt(10.5)
run.font.color.rgb = MUTED

doc.add_paragraph()

# ---------------------------------------------------------------- Intro
body(
    "WattWise Classrooms is a web application that helps a senior high school "
    "record how long classroom appliances run, estimate the energy and cost "
    "those appliances consume, and review consumption across rooms, sections, "
    "and months. It is a React frontend with a FastAPI backend, deployed on "
    "Vercel."
)
body(
    "The system was built through six phases. Each phase is summarised below "
    "with the work it produced in this project."
)

# ---------------------------------------------------------------- Flow diagram
heading("The Six Steps", 1)

steps = [
    ("1. PLANNING", "Define the problem, objectives, and scope"),
    ("2. REQUIREMENTS ANALYSIS", "Gather and specify what the system must do"),
    ("3. SYSTEM DESIGN", "Plan the architecture, data, and interfaces"),
    ("4. DEVELOPMENT", "Build the frontend, backend, and database"),
    ("5. TESTING", "Verify correctness, security, and usability"),
    ("6. DEPLOYMENT & MAINTENANCE", "Release the system and keep improving it"),
]

flow = doc.add_table(rows=len(steps), cols=1)
flow.style = "Table Grid"
for index, (name, detail) in enumerate(steps):
    cell = flow.cell(index, 0)
    shade(cell, "EAF3E9" if index % 2 == 0 else "F6FAF5")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name + "  —  ")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = ACCENT
    run2 = p.add_run(detail)
    run2.font.size = Pt(9.5)
caption("Figure 1. The six-step development process.")

# ---------------------------------------------------------------- Summary table
table(
    ["Step", "What it means", "What was done in WattWise"],
    [
        [
            "1. Planning",
            "Identify the problem and decide whether a solution is worth building.",
            "Confirmed that the school has no systematic way to monitor classroom "
            "electricity. Set objectives and checked technical, operational, and "
            "cost feasibility. Chose to build a browser-based estimator with no "
            "metering hardware.",
        ],
        [
            "2. Requirements Analysis",
            "Determine exactly what the system must do and for whom.",
            "Defined two users — teachers and administrators. Specified appliance "
            "logging, a custom-time entry for missed sessions, live energy and "
            "cost estimates, account issuance and approval, room and section "
            "assignment, and filtered consumption reports.",
        ],
        [
            "3. System Design",
            "Plan the structure before writing code.",
            "Designed a three-tier system: a React interface, a FastAPI service "
            "holding all business rules, and a relational database of five "
            "entities. Chose the E = P × t ÷ 1000 energy formula as the "
            "computational core.",
        ],
        [
            "4. Development",
            "Build the planned system.",
            "Implemented 32 API endpoints and the full user and administrator "
            "interfaces. Added live timers, preset appliance ratings, quantity "
            "support for multiple identical units, reporting with filters, and "
            "light and dark themes.",
        ],
        [
            "5. Testing",
            "Check that the system works and is safe to use.",
            "Ran lint and production builds, byte-compiled the backend, and "
            "exercised the API against the live deployment. Verified energy "
            "figures by hand, confirmed unauthorized access is rejected, and "
            "fixed five defects found during review.",
        ],
        [
            "6. Deployment & Maintenance",
            "Release the system and keep it current.",
            "Deployed the frontend and API to Vercel as production services with "
            "environment-based configuration. Subsequent maintenance updated "
            "appliance ratings, refined reporting, added interface animations, "
            "and reset demonstration data via a transactional maintenance tool.",
        ],
    ],
    widths=[1.2, 1.7, 3.6],
    font_size=8.5,
)
caption("Table 1. The six steps and their application to WattWise.")

doc.add_page_break()

# ---------------------------------------------------------------- Detail
heading("Step 1 — Planning", 1)
body(
    "The project began with a simple observation: the school had no reliable "
    "record of how much electricity each classroom used. Appliances were "
    "switched on and off without tracking, so no one could tell which rooms "
    "consumed the most, or what that consumption cost."
)
body(
    "Planning established the objectives — make logging simple, convert run "
    "time into energy and cost automatically, and give administrators a clear "
    "overview — and confirmed the project was feasible. Because no sensors or "
    "meters would be needed, the largest cost and complexity barriers were "
    "removed before development began."
)

heading("Step 2 — Requirements Analysis", 1)
body(
    "Two distinct users were identified, each with different needs:"
)
for item in [
    "Teachers need to log appliance use quickly, including sessions they forgot "
    "to time, and see an immediate estimate of energy and cost.",
    "Administrators need to issue accounts, approve registrations, assign "
    "rooms and sections, maintain the appliance catalog, and review "
    "consumption across the campus.",
]:
    p = doc.add_paragraph(item, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)

body(
    "The requirements also fixed an important constraint: the system produces "
    "estimates, not measurements. Every screen that shows a figure states this "
    "plainly, so estimates are never mistaken for meter readings."
)

heading("Step 3 — System Design", 1)
body(
    "The system was designed in three layers. The interface handles "
    "presentation only. The application layer holds every business rule, "
    "including the energy computation and all access control, so the browser "
    "can never become a competing source of truth. The database stores five "
    "entities: classroom sections, users, rooms, appliances, and usage records."
)
formula = doc.add_paragraph()
formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = formula.add_run("E = P × t ÷ 1000            E_total = Σ (Pᵢ × tᵢ) ÷ 1000")
run.bold = True
run.font.size = Pt(11.5)
formula.paragraph_format.space_after = Pt(8)
body(
    "Two design decisions proved important. First, rounding is applied once, "
    "so a displayed kWh figure multiplied by the tariff always equals the "
    "displayed bill. Second, an appliance can represent several identical "
    "units, so two fans are recorded and priced as one entry with a combined "
    "load rather than understating consumption."
)

heading("Step 4 — Development", 1)
body(
    "The interface and API were built in parallel. On the user side: a timed "
    "on/off control per appliance, a custom-time form for missed sessions with "
    "its own date, preset appliance ratings, and a live formula preview. On "
    "the administrator side: a dashboard, filtered reports, a 30-day room map "
    "that flags high consumption, and a merged console for managing users and "
    "facilities."
)
body(
    "Development followed an iterative approach. When the school refined its "
    "appliance ratings, the new values were applied and every derived figure "
    "recalculated. This flexibility is the main reason iteration was chosen "
    "over a single fixed plan."
)

heading("Step 5 — Testing", 1)
body(
    "Testing combined automated checks with manual and API verification. The "
    "frontend was linted and production-built; the backend was byte-compiled "
    "and exercised through its endpoints. Calculations were checked by hand — "
    "for example, 2 fans at 55 W, one TV at 100 W, and a 500 W dispenser total "
    "710 W, which is 5.68 kWh over eight hours."
)
table(
    ["Check", "Result"],
    [
        ["Energy and cost calculations", "Verified against hand computation"],
        ["Displayed kWh × tariff equals displayed bill", "Confirmed exact"],
        ["Unauthorized admin access", "Rejected"],
        ["Deleting an appliance with usage history", "Rejected (audit trail preserved)"],
        ["Entered hours beyond the limit", "Clamped to a 24-hour ceiling"],
        ["Light and dark themes across all routes", "Readable in both"],
    ],
    widths=[4.0, 2.5],
)
caption("Table 2. Principal verification outcomes.")
body(
    "Five defects were found and corrected, including a page that failed to "
    "render and a bill that disagreed with its own displayed kWh. Notably, "
    "the rendering fault passed both lint and build because the code was "
    "syntactically valid — it only appeared when the page was loaded. The "
    "project has no automated test suite, and this is the clearest gap to "
    "address next."
)

heading("Step 6 — Deployment & Maintenance", 1)
body(
    "The system was deployed to Vercel as two services: the static frontend "
    "and the Python API, with routing that separates API traffic from "
    "application pages. Credentials, the session secret, and the electricity "
    "tariff are supplied through environment variables so they can be changed "
    "without editing code."
)
body(
    "Maintenance has continued since release. Appliance ratings were revised, "
    "reporting gained filters and a search box for questions such as which "
    "room consumes the most, interface animations were added, and a "
    "transactional maintenance tool was built to back up and reset "
    "demonstration data safely."
)

heading("Conclusion", 1)
body(
    "WattWise Classrooms was delivered through six phases, from identifying the "
    "school's lack of energy visibility to a deployed, running system. The "
    "result lets a school record classroom appliance use, convert it into "
    "verifiable energy and cost estimates, and analyse consumption over time — "
    "without purchasing any metering hardware. The recommended next step is an "
    "automated test suite covering the energy model and access control."
)

doc.save("SDLC_WattWise_6Steps.docx")
print("Generated SDLC_WattWise_6Steps.docx")
