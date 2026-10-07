import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_flow_box(doc, text, fill="4A90E2"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.allow_autofit = False
    tbl.columns[0].width = Inches(4)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(255, 255, 255)
    return tbl

doc = docx.Document()
doc.add_paragraph('Figure 1: WattWise MCS Conceptual Framework').bold = True
p = doc.add_paragraph('SYSTEM FRAMEWORK')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold = True

boxes = [
    'System Initialization (Seeding & Schema)',
    'User/Admin Role Selection',
    'Credential & Room Assignment (Admin)',
    'Appliance Usage Tracking (User Timer)',
    'Energy Calculation (E=P*t/1000)',
    'Submission to API'
]

for b in boxes:
    create_flow_box(doc, b)
    p = doc.add_paragraph('↓')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Parameters table
table_out = doc.add_table(rows=1, cols=4)
headers = ['Usage Data\n(Appliance, Time)', 'Energy\nCalculation\n(kWh)', 'Room\nAggregation\n(30 days)', 'Admin Dashboard\nReports & Insights']
for i, text in enumerate(headers):
    cell = table_out.cell(0, i)
    set_cell_background(cell, '357ABD')
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(255, 255, 255)

doc.add_paragraph('↓').alignment = WD_ALIGN_PARAGRAPH.CENTER
create_flow_box(doc, 'Data Analysis & High-Use Detection', fill='2C3E50')
doc.add_paragraph('WattWise MCS Conceptual Framework').alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()
doc.add_heading('WattWise MCS Conceptual Framework', level=1)
doc.add_paragraph('The WattWise Classrooms (MCS) conceptual framework is designed to facilitate real-time electricity monitoring in educational institutions. The process begins with system initialization, where the SQLite database is automatically seeded with room, section, and appliance catalog data. Administrators manage access and room assignments via the admin console.')
doc.add_paragraph('Teachers use the User Workspace to log appliance runtimes using local browser timers. These submissions are processed by the FastAPI backend, which calculates electricity consumption using the formula: E = P × t / 1000 (kWh). Total consumption is aggregated across rooms for 30-day reporting cycles, identifying "high-use" classrooms.')
doc.add_paragraph('The Admin Portal utilizes this data for mapping usage trends, generating cost reports, and identifying anomalies (e.g., unusual duration of appliance use). This framework optimizes emergency and operational planning by providing transparent data-driven insights into classroom energy efficiency.')

doc.save('WattWise_Framework.docx')
