import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_flow_box(doc, text, fill="4A90E2"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.allow_autofit = False
    tbl.columns[0].width = Inches(4)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(255, 255, 255)
    return tbl

doc = docx.Document()
doc.add_paragraph('Figure 1').bold = True
p = doc.add_paragraph('FRAMEWORK')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold = True

boxes = [
    'Preparation of Materials',
    'Construction of Device',
    'Testing of the Device',
    'Applying ArcGIS and Google Maps Road Network Routing',
    'Deploying Device in Flood Warning Poles',
    'Training Data'
]

for b in boxes:
    create_flow_box(doc, b)
    p = doc.add_paragraph('↓')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Parameters table
table_out = doc.add_table(rows=1, cols=6)
headers = ['Location\nBarangay\n/Street', 'Water\nLevel\n(ft)', 'Alert Level\nSystem via\nTelegram', 'Echo\nDuration\nof\nBuzzer\nper trials', 'Flood Sensor\nStatus\n(color\nindicator)', 'Speed of real time\nnotification\nupdates via\nTelegram']
for i, text in enumerate(headers):
    cell = table_out.cell(0, i)
    set_cell_background(cell, '357ABD')
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(255, 255, 255)

doc.add_paragraph('↓').alignment = WD_ALIGN_PARAGRAPH.CENTER
create_flow_box(doc, 'Data Analysis', fill='2C3E50')
doc.add_paragraph('Conceptual Framework').alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()
doc.add_heading('Conceptual Framework', level=1)
doc.add_paragraph('This analysis outlines the development and assessment of FlogIS, an ArcGIS-based flood warning system integrated with Google Maps. The process began with material and device preparation, including the collection of GIS road network data by gathering the information about the flood hazard maps of the specific area and the selection of appropriate sensors and microcontrollers, and consultation with professionals for technical guidance.')
doc.add_paragraph('Following this, the flood warning device was assembled, configured with a Telegram notification feature, and integrated with ArcGIS and Google Maps routing for emergency response planning. Once the system was developed, testing and data collection were conducted, including simulation of rising water levels and calibration of alert thresholds.')
doc.add_paragraph('The device was then deployed in strategic locations based on the most affected area by the flood and the information given by the historical records of the place and connected to the routing system to enable real-time emergency planning. During the evaluation stage, the system\'s flood detection accuracy, Telegram alert speed and coverage, and routing efficiency were measured.')
doc.add_paragraph('Lastly, actual and predicted response times were compared, and the accuracy of sensor-based alerts was assessed and data analysis. This framework demonstrates how integrating sensors, GIS, and Telegram notification technology can enhance disaster preparedness, optimize emergency response, and improve community safety.')

doc.save('Conceptual_Framework.docx')
